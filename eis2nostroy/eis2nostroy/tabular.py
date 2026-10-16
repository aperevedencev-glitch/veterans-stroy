"""Чтение таблиц (XLSX/CSV) с гибким поиском столбцов по названию."""
from __future__ import annotations

import csv
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def norm(s: str) -> str:
    return re.sub(r"[^a-zа-яё0-9]+", "", str(s or "").lower())


def read_rows(path: Path) -> list[dict]:
    """Возвращает строки как словари {нормализованный_заголовок: значение}.

    Заголовок ищется в первых 10 строках: первая строка, где есть «инн».
    """
    path = Path(path)
    if path.suffix.lower() in (".csv", ".txt"):
        with open(path, encoding="utf-8-sig", newline="") as fh:
            sample = fh.read(4096)
            fh.seek(0)
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            rows = list(csv.reader(fh, dialect))
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header_idx = next(
        (i for i, r in enumerate(rows[:10]) if any("инн" in norm(c) for c in r if c)), 0
    )
    header = [norm(c) for c in rows[header_idx]]
    out = []
    for r in rows[header_idx + 1:]:
        if not any(v not in (None, "") for v in r):
            continue
        out.append({h: v for h, v in zip(header, r) if h})
    return out


def pick(row: dict, *candidates: str):
    """Значение первого столбца, название которого содержит один из вариантов."""
    for cand in candidates:
        c = norm(cand)
        for k, v in row.items():
            if c and c in k and v not in (None, ""):
                return v
    return None


def to_inn(v) -> str:
    if v is None:
        return ""
    s = re.sub(r"\D", "", str(v).split(".")[0])
    return s


def to_date(v) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:19]
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(s[: len(datetime.now().strftime(fmt))], fmt).date()
        except ValueError:
            continue
    return None


def to_float(v) -> float | None:
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v)).replace(",", ".")
    if s.count(".") > 1:
        head, _, tail = s.rpartition(".")
        s = head.replace(".", "") + "." + tail
    try:
        return float(s)
    except ValueError:
        return None
