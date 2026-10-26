"""Выгрузки: Excel для оператора и пакет для передачи в НОСТРОЙ."""
from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import (ACTIONS, CONFIRMED, DECLARED, IN_NOSTROY, MISMATCH, OUT_OF_PERIOD,
                     PROBABLE, ReconciledRecord)

LAW_NAMES = {"44": "44-ФЗ", "223": "223-ФЗ", "615": "ПП РФ 615", "": "не указан"}
FILL = {CONFIRMED: "D7EFDD", IN_NOSTROY: "D7EFDD", PROBABLE: "FCEBC8", DECLARED: "FCEBC8",
        MISMATCH: "F8D3CF", OUT_OF_PERIOD: "EEEEEE"}

# Поля пакета для НОСТРОЙ. Состав сверить со спецификацией НОСТРОЙ,
# когда она будет получена; до этого пакет загружает оператор.
NOSTROY_FIELDS = [
    ("member_inn", "ИНН члена СРО"),
    ("member_name", "Наименование члена СРО"),
    ("law", "Основание (закон)"),
    ("competitive", "Конкурентный способ"),
    ("reg_number", "Реестровый номер договора в ЕИС"),
    ("notice_number", "Номер извещения"),
    ("customer_inn", "ИНН заказчика"),
    ("customer_name", "Заказчик"),
    ("sign_date", "Дата заключения"),
    ("price", "Цена договора, ₽"),
    ("subject", "Предмет договора"),
    ("ikz", "ИКЗ (44-ФЗ)"),
    ("grk", "Градостроительный кодекс"),
    ("grk_note", "Основание отнесения к ГрК"),
    ("sources", "Источники"),
]


def _row(r: ReconciledRecord) -> dict:
    c = r.primary
    return {
        "member_inn": r.member.inn,
        "member_name": r.member.name,
        "law": LAW_NAMES.get(c.law, c.law),
        "competitive": {True: "да", False: "нет", None: ""}[c.competitive],
        "reg_number": c.reg_number,
        "notice_number": c.notice_number,
        "customer_inn": c.customer_inn,
        "customer_name": c.customer_name,
        "sign_date": c.sign_date.strftime("%d.%m.%Y") if c.sign_date else "",
        "price": c.price,
        "subject": c.subject,
        "ikz": c.extra.get("ikz", ""),
        "grk": c.extra.get("grk", ""),
        "grk_note": c.extra.get("grk_note", ""),
        "sources": ", ".join(r.sources),
    }


def write_operator_workbook(records: list[ReconciledRecord], path: Path, sro_name: str = "") -> Path:
    wb = Workbook()
    thin = Side(style="thin", color="C9CED6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F3A5F")

    def sheet(ws, title, items):
        ws.append([title])
        ws["A1"].font = Font(name="Arial", size=13, bold=True, color="1F3A5F")
        ws.append([f"Сформировано {datetime.now():%d.%m.%Y %H:%M}. Записей: {len(items)}"])
        ws.append([])
        cols = ["Статус", "Действие оператора"] + [h for _, h in NOSTROY_FIELDS] + ["Замечания"]
        widths = [16, 30, 14, 34, 10, 10, 24, 22, 14, 36, 12, 16, 40, 38, 14, 40, 18, 40]
        ws.append(cols)
        for j, w in enumerate(widths, 1):
            cell = ws.cell(4, j)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = w
        for r in items:
            d = _row(r)
            ws.append([r.status, r.action] + [d[k] for k, _ in NOSTROY_FIELDS] + ["; ".join(r.issues)])
            row = ws.max_row
            for j in range(1, len(cols) + 1):
                cell = ws.cell(row, j)
                cell.font = Font(name="Arial", size=10)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=j in (2, 6, 10, 13, 15))
            ws.cell(row, 1).fill = PatternFill("solid", fgColor=FILL.get(r.status, "FFFFFF"))
            ws.cell(row, 12).number_format = "#,##0.00"
        ws.freeze_panes = "C5"
        if items:
            ws.auto_filter.ref = f"A4:{ws.cell(4, len(cols)).column_letter}{ws.max_row}"

    summary = wb.active
    summary.title = "Сводка"
    summary.append([f"Сверка договоров членов СРО {sro_name}".strip()])
    summary["A1"].font = Font(name="Arial", size=14, bold=True, color="1F3A5F")
    summary.append([f"Сформировано {datetime.now():%d.%m.%Y %H:%M}"])
    summary.append([])
    summary.append(["Статус", "Записей", "Действие"])
    counts = Counter(r.status for r in records)
    for st in (CONFIRMED, PROBABLE, MISMATCH, DECLARED, IN_NOSTROY, OUT_OF_PERIOD):
        summary.append([st, counts.get(st, 0), ACTIONS[st]])
        summary.cell(summary.max_row, 1).fill = PatternFill("solid", fgColor=FILL[st])
    summary.append(["Всего", len(records), ""])
    summary.cell(summary.max_row, 1).font = Font(name="Arial", bold=True)
    for col, w in (("A", 22), ("B", 10), ("C", 36)):
        summary.column_dimensions[col].width = w
    for j in range(1, 4):
        summary.cell(4, j).font = Font(name="Arial", bold=True, color="FFFFFF")
        summary.cell(4, j).fill = head_fill

    sheet(wb.create_sheet("К передаче"), "К передаче в НОСТРОЙ",
          [r for r in records if r.status == CONFIRMED])
    sheet(wb.create_sheet("Требуют решения"), "Требуют решения оператора",
          [r for r in records if r.status in (PROBABLE, MISMATCH, DECLARED)])
    sheet(wb.create_sheet("Все записи"), "Все записи", records)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_nostroy_package(records: list[ReconciledRecord], out_dir: Path) -> tuple[Path, Path]:
    """CSV и JSON только с подтверждёнными записями, готовыми к передаче."""
    items = [_row(r) for r in records if r.status == CONFIRMED]
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    csv_path = out_dir / f"nostroy_{stamp}.csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow([h for _, h in NOSTROY_FIELDS])
        for d in items:
            w.writerow([d[k] if d[k] is not None else "" for k, _ in NOSTROY_FIELDS])
    json_path = out_dir / f"nostroy_{stamp}.json"
    json_path.write_text(json.dumps({"generated": datetime.now().isoformat(timespec="seconds"),
                                     "count": len(items), "items": items},
                                    ensure_ascii=False, indent=2), encoding="utf-8")
    return csv_path, json_path
