"""Тесты на синтетических данных (ИНН и договоры вымышленные)."""
from __future__ import annotations

import io
import json
import zipfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from eis2nostroy import cli
from eis2nostroy.config import load_config
from eis2nostroy.eis_client import EisClient, EisError, parse_archive_urls
from eis2nostroy.models import (CONFIRMED, DECLARED, IN_NOSTROY, MISMATCH, OUT_OF_PERIOD,
                                PROBABLE, Contract, Member)
from eis2nostroy.parser import parse_archive, parse_contract_xml
from eis2nostroy.reconcile import reconcile

MEMBER_A = "7700000001"   # член с уровнем ОДО с 2019 года
MEMBER_B = "5000000002"   # член без даты ОДО
OUTSIDER = "7799999999"

XML_44 = f"""<?xml version="1.0" encoding="UTF-8"?>
<export xmlns="http://zakupki.gov.ru/oos/export/1" xmlns:oos="http://zakupki.gov.ru/oos/types/1">
 <contract>
  <oos:regNum>2770100000126000001</oos:regNum>
  <oos:signDate>2026-05-12</oos:signDate>
  <oos:customer><oos:regNum>01731000001</oos:regNum><oos:fullName>ГКУ Тестовый заказчик</oos:fullName><oos:inn>7701000001</oos:inn></oos:customer>
  <oos:foundation><oos:fcsOrder><oos:notificationNumber>0373200000126000111</oos:notificationNumber>
    <oos:placingWay><oos:name>Электронный аукцион</oos:name></oos:placingWay></oos:fcsOrder></oos:foundation>
  <oos:priceInfo><oos:price>12500000.50</oos:price></oos:priceInfo>
  <oos:contractSubject>Капитальный ремонт кровли</oos:contractSubject>
  <oos:suppliers><oos:supplier><oos:legalEntityRF><oos:fullName>ООО Пример</oos:fullName><oos:INN>{MEMBER_A}</oos:INN></oos:legalEntityRF></oos:supplier></oos:suppliers>
 </contract>
</export>"""

XML_44_OTHER = XML_44.replace(MEMBER_A, OUTSIDER).replace("2770100000126000001", "2770100000126000002")

XML_223 = f"""<?xml version="1.0" encoding="UTF-8"?>
<ns2:contract xmlns:ns2="http://zakupki.gov.ru/223fz/contract/1">
 <ns2:body><ns2:item><ns2:contractData>
  <ns2:registrationNumber>57701000001260000010000</ns2:registrationNumber>
  <ns2:contractDate>2026-08-01</ns2:contractDate>
  <ns2:customer><ns2:mainInfo><ns2:fullName>АО Тестовая сетевая компания</ns2:fullName><ns2:inn>7702000002</ns2:inn></ns2:mainInfo></ns2:customer>
  <ns2:supplierInfo><ns2:inn>{MEMBER_B}</ns2:inn><ns2:name>ООО Подрядчик Б</ns2:name></ns2:supplierInfo>
  <ns2:purchaseMethodName>Закупка у единственного поставщика</ns2:purchaseMethodName>
  <ns2:sum>3000000</ns2:sum>
  <ns2:subject>Ремонт подстанции</ns2:subject>
 </ns2:contractData></ns2:item></ns2:body>
</ns2:contract>"""


def make_zip(path: Path, files: dict[str, str], nested: bool = False) -> Path:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, text in files.items():
            zf.writestr(name, text)
    if nested:
        outer = io.BytesIO()
        with zipfile.ZipFile(outer, "w") as zf:
            zf.writestr("inner.zip", buf.getvalue())
        buf = outer
    path.write_bytes(buf.getvalue())
    return path


# ---------------------------------------------------------------- разбор XML
def test_parse_44_finds_member_only():
    got = parse_contract_xml(XML_44.encode(), "44", {MEMBER_A, MEMBER_B})
    assert len(got) == 1
    c = got[0]
    assert c.supplier_inn == MEMBER_A
    assert c.reg_number == "2770100000126000001"
    assert c.customer_inn == "7701000001"
    assert c.sign_date == date(2026, 5, 12)
    assert c.price == pytest.approx(12500000.50)
    assert c.notice_number == "0373200000126000111"
    assert c.competitive is True
    assert parse_contract_xml(XML_44_OTHER.encode(), "44", {MEMBER_A}) == []


def test_parse_223_noncompetitive():
    c = parse_contract_xml(XML_223.encode(), "223", {MEMBER_B})[0]
    assert c.reg_number == "57701000001260000010000"
    assert c.customer_inn == "7702000002"
    assert c.price == 3000000
    assert c.competitive is False
    assert c.sign_date == date(2026, 8, 1)


def test_parse_nested_archive(tmp_path):
    z = make_zip(tmp_path / "a.zip", {"c1.xml": XML_44, "c2.xml": XML_44_OTHER, "bad.xml": "<oops"}, nested=True)
    got = parse_archive(z, "44", {MEMBER_A})
    assert [c.supplier_inn for c in got] == [MEMBER_A]


# ---------------------------------------------------------------- клиент ЕИС
class FakeResp:
    def __init__(self, status, content):
        self.status_code, self.content = status, content
        self.text = content.decode("utf-8", "ignore") if isinstance(content, bytes) else content


class FakeSession:
    def __init__(self, archive: bytes):
        self.posts, self.gets, self.archive = [], [], archive

    def post(self, url, data, headers, timeout):
        self.posts.append((url, data.decode()))
        return FakeResp(200, b"""<soap:Envelope xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/"><soap:Body>
<ns2:getDocsByOrgRegionResponse xmlns:ns2="http://zakupki.gov.ru/fz44/get-docs-ip/ws"><dataInfo>
<archiveUrl>https://int44.zakupki.gov.ru/archive/1.zip</archiveUrl></dataInfo></ns2:getDocsByOrgRegionResponse></soap:Body></soap:Envelope>""")

    def get(self, url, headers, timeout):
        self.gets.append((url, headers))
        # первая попытка — архив ещё не готов
        if len(self.gets) == 1:
            return FakeResp(404, b"not ready")
        return FakeResp(200, self.archive)


def write_config(tmp_path: Path, extra: str = "") -> Path:
    src = (Path(__file__).parents[1] / "config.example.toml").read_text(encoding="utf-8")
    src = src.replace('token = ""', 'token = "TEST-TOKEN"', 1)
    src = src.replace("min_interval_sec = 1.5", "min_interval_sec = 0")
    src = src.replace("archive_retry_wait_sec = 20", "archive_retry_wait_sec = 0")
    p = tmp_path / "config.toml"
    p.write_text(src + extra, encoding="utf-8")
    return p


def test_client_request_and_download(tmp_path):
    cfg = load_config(write_config(tmp_path))
    archive = make_zip(tmp_path / "src.zip", {"c.xml": XML_44}).read_bytes()
    sess = FakeSession(archive)
    client = EisClient(cfg, session=sess)
    paths = client.fetch_day(cfg.feeds[0], "77", date(2026, 5, 12))
    assert len(paths) == 1 and paths[0].read_bytes()[:2] == b"PK"
    url, body = sess.posts[0]
    assert url.endswith("/getDocsLE")
    assert "<subsystemType>RGK</subsystemType>" in body
    assert "<orgRegion>77</orgRegion>" in body
    assert "<exactDate>2026-05-12</exactDate>" in body
    assert "TEST-TOKEN" in body
    assert sess.gets[-1][1] == {"individualPerson_token": "TEST-TOKEN"}
    # повторный вызов берёт кэш и не обращается к ЕИС
    client.fetch_day(cfg.feeds[0], "77", date(2026, 5, 12))
    assert len(sess.posts) == 1


def test_soap_fault_raises():
    with pytest.raises(EisError):
        parse_archive_urls("""<soap:Envelope xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/"><soap:Body>
<soap:Fault><faultstring>Токен недействителен</faultstring></soap:Fault></soap:Body></soap:Envelope>""".encode())


# ---------------------------------------------------------------- сверка
def _c(source, **kw):
    base = dict(source=source, law="44", supplier_inn=MEMBER_A, customer_inn="7701000001",
                sign_date=date(2026, 5, 12), price=1_000_000.0, competitive=True)
    base.update(kw)
    return Contract(**base)


def test_statuses():
    members = {MEMBER_A: Member(MEMBER_A, "А", odo_date=date(2019, 5, 15)),
               MEMBER_B: Member(MEMBER_B, "Б")}
    contracts = [
        _c("eis_44", reg_number="R1"),                                      # подтверждён
        _c("eis_44", reg_number="R2", sign_date=date(2026, 6, 1)),
        _c("nostroy", reg_number="R2", sign_date=date(2026, 6, 1)),         # уже в НОСТРОЙ
        _c("eis_44", reg_number="R3", sign_date=date(2026, 7, 1)),
        _c("finpod", reg_number="R3", sign_date=date(2026, 7, 1), price=1_500_000.0),  # расхождение
        _c("finpod", sign_date=date(2026, 8, 1), price=2_000_000.0),        # вероятный
        _c("notice", sign_date=date(2026, 9, 1), price=3_000_000.0),        # заявлен членом
        _c("eis_44", reg_number="R4", sign_date=date(2018, 1, 10)),         # до уровня ОДО
        _c("eis_223", law="223", reg_number="R5", supplier_inn=MEMBER_B,
           sign_date=date(2025, 12, 1), competitive=False),                  # неконкурентный до 01.03.2026
        _c("eis_44", reg_number="X", supplier_inn=OUTSIDER),                 # не член СРО
    ]
    recs = reconcile(members, contracts, date(2017, 7, 1), date(2026, 3, 1))
    by_reg = {r.primary.reg_number or r.primary.source: r.status for r in recs}
    assert by_reg["R1"] == CONFIRMED
    assert by_reg["R2"] == IN_NOSTROY
    assert by_reg["R3"] == MISMATCH
    assert by_reg["finpod"] == PROBABLE
    assert by_reg["notice"] == DECLARED
    assert by_reg["R4"] == OUT_OF_PERIOD
    assert by_reg["R5"] == OUT_OF_PERIOD
    assert "X" not in by_reg
    assert len(recs) == 7


def test_fuzzy_match_without_number():
    members = {MEMBER_A: Member(MEMBER_A)}
    recs = reconcile(members, [
        _c("eis_44", reg_number="R1"),
        _c("finpod", sign_date=date(2026, 5, 14), price=1_005_000.0),       # ±2 дня, ±0,5 %
    ], date(2017, 7, 1), date(2026, 3, 1))
    assert len(recs) == 1 and recs[0].status == CONFIRMED
    assert recs[0].sources == ["eis_44", "finpod"]


# ---------------------------------------------------------------- сквозной прогон
def test_end_to_end(tmp_path, monkeypatch):
    data = tmp_path / "data"
    data.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["Члены СРО"])
    ws.append([])
    ws.append(["ИНН члена", "Наименование", "Дата присвоения уровня по ОДО"])
    ws.append([MEMBER_A, "ООО Пример", "15.05.2019"])
    ws.append([MEMBER_B, "ООО Подрядчик Б", None])
    wb.save(data / "members.xlsx")
    (data / "finpod.csv").write_text(
        "ИНН поставщика;Номер извещения;ИНН заказчика;Дата договора;Сумма\n"
        f"{MEMBER_A};0373200000126000999;7703000003;20.09.2026;4 400 000,00\n", encoding="utf-8")
    (data / "nostroy.csv").write_text(
        "ИНН члена;Реестровый номер;Дата договора;Цена\n"
        f"{MEMBER_B};57701000001260000010000;01.08.2026;3000000\n", encoding="utf-8")
    cfg_path = write_config(tmp_path)
    text = cfg_path.read_text(encoding="utf-8")
    text = text.replace('finpod_import = ""', 'finpod_import = "data/finpod.csv"')
    text = text.replace('nostroy_registry = ""', 'nostroy_registry = "data/nostroy.csv"')
    cfg_path.write_text(text, encoding="utf-8")

    arch = data / "archives"
    arch.mkdir()
    make_zip(arch / "rgk.zip", {"c1.xml": XML_44, "c2.xml": XML_44_OTHER})
    assert cli.main(["-c", str(cfg_path), "import-archives", "--law", "44", str(arch)]) == 0
    arch2 = data / "archives223"
    arch2.mkdir()
    make_zip(arch2 / "rd223.zip", {"c.xml": XML_223})
    assert cli.main(["-c", str(cfg_path), "import-archives", "--law", "223", str(arch2)]) == 0
    assert cli.main(["-c", str(cfg_path), "reconcile"]) == 0

    out = data / "out"
    book = next(out.glob("sverka_*.xlsx"))
    wb = load_workbook(book)
    assert wb.sheetnames == ["Сводка", "К передаче", "Требуют решения", "Все записи"]
    to_send = [r for r in wb["К передаче"].iter_rows(min_row=5, values_only=True) if r[0]]
    assert len(to_send) == 1 and to_send[0][2] == MEMBER_A
    pending = [r for r in wb["Требуют решения"].iter_rows(min_row=5, values_only=True) if r[0]]
    assert [r[0] for r in pending] == [PROBABLE]
    js = json.loads(next(out.glob("nostroy_*.json")).read_text(encoding="utf-8"))
    assert js["count"] == 1 and js["items"][0]["reg_number"] == "2770100000126000001"
    # неконкурентный договор 223-ФЗ уже есть в реестре НОСТРОЙ
    allrows = [r for r in wb["Все записи"].iter_rows(min_row=5, values_only=True) if r[0]]
    assert sorted(r[0] for r in allrows) == sorted([CONFIRMED, PROBABLE, IN_NOSTROY])
